"""Convierte una hoja del dataset Online Retail II a CSV reproducible.

Uso:
    python scripts/prepare_historical.py \
        --xlsx "ruta/online_retail_II.xlsx" \
        --sheet "Year 2010-2011" \
        --output "data/incoming/historical/online_retail_2010_2011.csv"

El archivo original NO se versiona. El CSV resultante tampoco: se genera con
este script y su hash queda en el metadata para trazabilidad.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

HEADERS = [
    "invoice_no",
    "stock_code",
    "description",
    "quantity",
    "invoice_date",
    "unit_price",
    "customer_id",
    "country",
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def as_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def as_int_text(value: object) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def as_float_text(value: object) -> str:
    if value is None or value == "":
        return ""
    return str(value)


def as_datetime_text(value: object) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day).isoformat(sep=" ")
    return str(value).strip()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", required=True, type=Path)
    parser.add_argument("--sheet", default="Year 2010-2011")
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument(
        "--limit",
        type=int,
        default=0,
        help="0 = todas las filas; util para una muestra rapida.",
    )
    args = parser.parse_args()

    if not args.xlsx.exists():
        raise SystemExit(f"No existe el archivo: {args.xlsx}")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    started_at = datetime.utcnow().isoformat(timespec="seconds") + "Z"

    workbook = load_workbook(args.xlsx, read_only=True, data_only=True)
    if args.sheet not in workbook.sheetnames:
        raise SystemExit(
            f"Hoja no encontrada: {args.sheet}. Disponibles: {workbook.sheetnames}"
        )
    sheet = workbook[args.sheet]

    rows_written = 0
    rows_seen = 0
    with args.output.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(HEADERS)
        for row in sheet.iter_rows(min_row=2, values_only=True):
            rows_seen += 1
            if row is None or all(value is None for value in row):
                continue
            invoice, stock, description, quantity, invoice_date, price, customer, country = (
                list(row) + [None] * 8
            )[:8]
            writer.writerow(
                [
                    as_text(invoice),
                    as_text(stock),
                    as_text(description),
                    as_int_text(quantity),
                    as_datetime_text(invoice_date),
                    as_float_text(price),
                    as_int_text(customer),
                    as_text(country),
                ]
            )
            rows_written += 1
            if args.limit and rows_written >= args.limit:
                break

    metadata = {
        "source_file": args.xlsx.name,
        "source_sha256": sha256(args.xlsx),
        "source_url": "https://archive.ics.uci.edu/dataset/502/online+retail+ii",
        "source_reference": "https://www.kaggle.com/datasets/thedevastator/online-retail-transaction-dataset",
        "license_note": "Ver licenses/README en UCI/Kaggle. Uso academico.",
        "sheet": args.sheet,
        "rows_seen": rows_seen,
        "rows_written": rows_written,
        "started_at_utc": started_at,
        "finished_at_utc": datetime.utcnow().isoformat(timespec="seconds") + "Z",
        "output_file": args.output.name,
        "output_sha256": sha256(args.output),
    }
    metadata_path = args.output.with_suffix(".metadata.json")
    metadata_path.write_text(
        json.dumps(metadata, indent=2, ensure_ascii=False), encoding="utf-8"
    )

    print(f"Filas escritas: {rows_written:,}")
    print(f"CSV: {args.output}")
    print(f"Metadata: {metadata_path}")


if __name__ == "__main__":
    main()
